"""Tests for BinaryReadGuardMiddleware."""

from __future__ import annotations

from pathlib import Path

from octop.infra.agents.middleware.binary_read_guard import (
    preview_spreadsheet,
    read_file_block_reason,
)


def test_blocks_inbound_pdf() -> None:
    reason = read_file_block_reason("inbound/01KWB9MG2Z570P7QB367KDJ75R.pdf")
    assert reason is not None
    assert "pdf" in reason.lower()
    assert "execute_shell_command" in reason


def test_allows_inbound_txt() -> None:
    assert read_file_block_reason("inbound/01JTEST.txt") is None


def test_blocks_pdf_anywhere() -> None:
    assert read_file_block_reason("docs/report.pdf") is not None


def test_does_not_use_workspace_for_pc_desktop_xlsx(tmp_path: Path) -> None:
    from octop.infra.agents.middleware.binary_read_guard import BinaryReadGuardMiddleware

    desk = tmp_path / "Desktop"
    desk.mkdir()
    (desk / "托管单位.xlsx").write_bytes(b"not-xlsx")
    mw = BinaryReadGuardMiddleware(workspace_dir=tmp_path)
    assert mw._preview_message(r"C:\Users\a\Desktop\托管单位.xlsx") is None
    assert mw._preview_message("/C:/Users/a/Desktop/托管单位.xlsx") is None
    assert mw._preview_message("Desktop/托管单位.xlsx") is None
    assert mw._preview_message(r"C:\Users\a\Desktop\合同.pdf") is None
    assert mw._preview_message(r"C:\Users\a\Desktop\shot.png") is None


def test_preview_xlsx(tmp_path: Path) -> None:
    from openpyxl import Workbook

    inbound = tmp_path / "inbound"
    inbound.mkdir()
    path = inbound / "units.xlsx"
    wb = Workbook()
    sheet = wb.active
    assert sheet is not None
    sheet.title = "单位"
    sheet.append(["名称", "账号"])
    sheet.append(["甲", "a1"])
    wb.save(path)
    text = preview_spreadsheet(path)
    assert "Spreadsheet: units.xlsx" in text
    assert "单位" in text
    assert "甲" in text


def test_preview_inbound_pdf_and_docx(tmp_path: Path) -> None:
    from docx import Document
    from pypdf import PdfWriter

    inbound = tmp_path / "inbound"
    inbound.mkdir()
    pdf_path = inbound / "note.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with pdf_path.open("wb") as handle:
        writer.write(handle)
    from octop.infra.agents.office_preview import dump_office

    pdf_text = dump_office(pdf_path)
    assert "PDF: note.pdf" in pdf_text

    doc_path = inbound / "letter.docx"
    document = Document()
    document.add_paragraph("托管单位甲")
    document.save(doc_path)
    from octop.infra.agents.middleware.binary_read_guard import BinaryReadGuardMiddleware

    mw = BinaryReadGuardMiddleware(workspace_dir=tmp_path)
    preview = mw._preview_message("inbound/letter.docx")
    assert preview is not None
    assert "托管单位甲" in preview
