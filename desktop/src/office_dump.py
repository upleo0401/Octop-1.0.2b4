"""Dump local Office/PDF files to UTF-8 text for the desktop Worker."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

MAX_CHARS = 24_000
MAX_ROWS = 80
MAX_COLS = 30


def clip(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS] + "\n... (truncated)"
    return text


def dump_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts = [f"Spreadsheet: {path.name}"]
    try:
        for name in wb.sheetnames:
            sheet = wb[name]
            parts.append(f"\n## Sheet: {name}")
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS:
                    parts.append("... (truncated rows)")
                    break
                cells = ["" if cell is None else str(cell) for cell in row[:MAX_COLS]]
                parts.append("\t".join(cells))
                if sum(len(item) for item in parts) > MAX_CHARS:
                    parts.append("... (truncated)")
                    return "\n".join(parts)
        return "\n".join(parts)
    finally:
        wb.close()


def dump_xls(path: Path) -> str:
    import xlrd

    book = xlrd.open_workbook(path)
    parts = [f"Spreadsheet: {path.name}"]
    for sheet in book.sheets():
        parts.append(f"\n## Sheet: {sheet.name}")
        rows = min(sheet.nrows, MAX_ROWS)
        for row_idx in range(rows):
            cells = [str(sheet.cell_value(row_idx, col)) for col in range(min(sheet.ncols, MAX_COLS))]
            parts.append("\t".join(cells))
        if sheet.nrows > MAX_ROWS:
            parts.append("... (truncated rows)")
        if sum(len(item) for item in parts) > MAX_CHARS:
            parts.append("... (truncated)")
            break
    return "\n".join(parts)


def dump_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [f"Document: {path.name}"]
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        parts.append("")
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))
    return clip("\n".join(parts))


def dump_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = [f"Presentation: {path.name}"]
    for index, slide in enumerate(prs.slides, start=1):
        parts.append(f"\n## Slide {index}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                table = shape.table
                for row in table.rows:
                    parts.append("\t".join(cell.text.strip() for cell in row.cells))
        if sum(len(item) for item in parts) > MAX_CHARS:
            parts.append("... (truncated)")
            break
    return clip("\n".join(parts))


def dump_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = [f"PDF: {path.name} ({len(reader.pages)} pages)"]
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        parts.append(f"\n## Page {index}")
        parts.append(text)
        if sum(len(item) for item in parts) > MAX_CHARS:
            parts.append("... (truncated)")
            break
    return clip("\n".join(parts))


def dump_via_office_com(path: Path, app: str) -> str:
    prog = {"word": "Word.Application", "powerpoint": "PowerPoint.Application"}[app]
    opener = {
        "word": (
            "$doc = $app.Documents.Open($path, $false, $true); "
            "$text = $doc.Content.Text; $doc.Close($false)"
        ),
        "powerpoint": (
            "$pres = $app.Presentations.Open($path, $true, $false, $false); "
            "$bits = @(); foreach ($s in $pres.Slides) { "
            "foreach ($sh in $s.Shapes) { if ($sh.HasTextFrame -eq -1) { "
            "$bits += $sh.TextFrame.TextRange.Text } } }; "
            "$text = $bits -join \"`n\"; $pres.Close()"
        ),
    }[app]
    script = (
        f"$ErrorActionPreference='Stop'; $path={_ps_quote(str(path))}; "
        f"$app=New-Object -ComObject {prog}; $app.Visible=$false; "
        f"{opener}; try {{ $app.Quit() }} catch {{}}; "
        "[Console]::OutputEncoding=[Text.UTF8Encoding]::new(); "
        "[Console]::Out.Write($text)"
    )
    result = subprocess.run(
        [
            "powershell.exe",
            "-NoProfile",
            "-WindowStyle",
            "Hidden",
            "-Command",
            script,
        ],
        check=False,
        capture_output=True,
        timeout=60,
    )
    if result.returncode != 0:
        err = (result.stderr or result.stdout).decode("utf-8", errors="replace").strip()
        raise RuntimeError(err or f"{app} COM dump failed")
    return clip(f"{path.name}\n" + result.stdout.decode("utf-8", errors="replace"))


def _ps_quote(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def dump(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in {".xlsx", ".xlsm"}:
        return dump_xlsx(path)
    if ext == ".xls":
        return dump_xls(path)
    if ext == ".docx":
        return dump_docx(path)
    if ext == ".pptx":
        return dump_pptx(path)
    if ext == ".pdf":
        return dump_pdf(path)
    if ext == ".doc":
        return dump_via_office_com(path, "word")
    if ext == ".ppt":
        return dump_via_office_com(path, "powerpoint")
    raise RuntimeError(f"unsupported type {ext}")


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("usage: office_dump.py <path>")
    path = Path(sys.argv[1])
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stdout.write(dump(path))


if __name__ == "__main__":
    main()
